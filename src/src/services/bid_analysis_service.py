"""Bid comparison and normalization utilities."""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO
from statistics import median
from typing import Any, Dict, Iterable, List, Optional, Sequence
import json

import numpy as np
import pandas as pd
from docx import Document

from services.llm_helpers import dataframe_to_pretty_json, run_agent_sync


SUPPLIER_ID_COLUMNS = [
    "supplier",
    "supplier_name",
    "vendor",
    "vendor_name",
    "company",
    "entity",
]

COLUMN_ALIASES: dict[str, list[str]] = {
    "supplier": SUPPLIER_ID_COLUMNS,
    "price": ["price", "total_price", "total_cost", "bid_total", "bid", "amount"],
    "quality": ["quality", "quality_score", "service_quality", "score_quality"],
    "lead_time": ["lead_time", "delivery_time", "leadtime", "sla_days", "delivery_days"],
    "warranty": ["warranty", "support", "support_months", "warranty_months"],
    "compliance": ["compliance", "compliance_score", "regulatory", "policy_alignment"],
    "risk": ["risk", "risk_score", "supplier_risk", "exposure"],
    "sustainability": ["sustainability", "esg", "sustainability_score", "carbon_score"],
    "hidden_costs": ["hidden_costs", "additional_fees", "surcharges", "extras"],
    "scope": ["scope_deviation", "scope_delta", "scope"],
}

NEGATIVE_METRICS = {"price", "lead_time", "risk", "hidden_costs"}

DEFAULT_METRIC_WEIGHTS = {
    "price": 0.25,
    "quality": 0.2,
    "lead_time": 0.15,
    "warranty": 0.1,
    "compliance": 0.15,
    "risk": 0.1,
    "sustainability": 0.05,
}


def _load_docx_document(data: bytes) -> Document:
    stream = BytesIO(data)
    stream.seek(0)
    return Document(stream)


def _docx_to_dataframe(data: bytes) -> pd.DataFrame:
    document = _load_docx_document(data)
    frames: List[pd.DataFrame] = []
    for table_idx, table in enumerate(document.tables):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if not rows:
            continue
        header = rows[0]
        data_rows = rows[1:] if len(rows) > 1 else []
        has_header = bool(data_rows) and any(header)
        if has_header and len(set(h.lower() for h in header if h)) == len([h for h in header if h]):
            columns = [h if h else f"column_{idx}" for idx, h in enumerate(header)]
        else:
            columns = [f"column_{idx}" for idx in range(len(header))]
            data_rows = rows
        frame = pd.DataFrame(data_rows, columns=columns)
        if not frame.empty:
            frames.append(frame)
    if frames:
        return pd.concat(frames, ignore_index=True)

    anchors = _docx_to_key_values(document)
    if anchors:
        return pd.DataFrame([anchors])
    raise ValueError("DOCX file does not contain tabular bid data.")


def _docx_to_key_values(document: Document) -> Dict[str, Any]:
    anchors: Dict[str, Any] = {}
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if len(cells) >= 2 and cells[0]:
                anchors[cells[0]] = cells[1]
    if anchors:
        return anchors
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text or ":" not in text:
            continue
        key, value = text.split(":", 1)
        anchors[key.strip()] = value.strip()
    return anchors


def _coerce_anchor_value(value: Any) -> Any:
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, str):
        stripped = value.strip()
        try:
            return float(stripped)
        except ValueError:
            return stripped
    return value


@dataclass
class WeightingConfig:
    raw_weights: Dict[str, float] = field(default_factory=lambda: DEFAULT_METRIC_WEIGHTS.copy())

    def normalized(self) -> Dict[str, float]:
        filtered = {k: float(v) for k, v in self.raw_weights.items() if float(v) >= 0}
        total = sum(filtered.values())
        if not filtered or total == 0:
            raise ValueError("At least one positive weight is required.")
        return {k: v / total for k, v in filtered.items()}


@dataclass
class SupplierScore:
    supplier: str
    total_score: float
    metrics: Dict[str, float]
    normalized_metrics: Dict[str, float]
    rationale: str
    risk_flags: List[str]


@dataclass
class BidComparisonResult:
    normalized_table: pd.DataFrame
    weight_config: Dict[str, float]
    supplier_scores: List[SupplierScore]
    deltas_vs_rfp: List[str]
    what_if_history: List[Dict[str, Any]]
    llm_summary: str
    generated_at: datetime = field(default_factory=datetime.utcnow)

    def to_csv(self) -> bytes:
        return self.normalized_table.to_csv(index=False).encode("utf-8")

    def to_markdown(self) -> str:
        return self.normalized_table.to_markdown(index=False)


@dataclass
class RFPScenarioScore:
    rfp_name: str
    total_score: float
    metrics: Dict[str, float]
    normalized_metrics: Dict[str, float]
    rationale: str
    fit_flags: List[str]


@dataclass
class MultiRFPComparisonResult:
    scenario_table: pd.DataFrame
    weight_config: Dict[str, float]
    rfp_scores: List[RFPScenarioScore]
    variation_notes: List[str]
    llm_summary: str
    generated_at: datetime = field(default_factory=datetime.utcnow)



class BidComparisonEngine:
    def __init__(
        self,
        *,
        include_outliers: bool,
        missing_strategy: str,
        weight_config: WeightingConfig,
        rfp_anchor: Optional[dict] = None,
    ) -> None:
        self.include_outliers = include_outliers
        self.missing_strategy = missing_strategy
        self.weight_config = weight_config
        self.rfp_anchor = rfp_anchor or {}
        self._what_if_history: List[Dict[str, Any]] = []
        self._baseline_df: Optional[pd.DataFrame] = None
        self._latest_scores: Dict[str, float] = {}

    def analyse(self, files: Sequence[Any]) -> BidComparisonResult:
        raw_frames = [self._load_file(f) for f in files if f]
        if not raw_frames:
            raise ValueError("Upload at least one supplier bid file.")

        merged = pd.concat(raw_frames, ignore_index=True)
        cleaned = self._prepare_dataset(merged)
        self._baseline_df = cleaned.copy()
        scored, supplier_scores = self._score_suppliers(cleaned)
        deltas = self._derive_deltas(scored)
        summary = self._llm_summary(scored, supplier_scores, deltas)
        self._latest_scores = {score.supplier: score.total_score for score in supplier_scores}
        return BidComparisonResult(
            normalized_table=scored,
            weight_config=self.weight_config.normalized(),
            supplier_scores=supplier_scores,
            deltas_vs_rfp=deltas,
            what_if_history=self._what_if_history.copy(),
            llm_summary=summary,
        )

    def run_scenario(self, weights: Dict[str, float]) -> Dict[str, Any]:
        scenario_config = WeightingConfig(weights)
        normalized_weights = scenario_config.normalized()
        baseline = getattr(self, "_baseline_df", None)
        if baseline is None:
            raise ValueError("Run primary analysis before executing scenarios.")
        recalculated, scores = self._score_suppliers(baseline.copy(), override_weights=normalized_weights)
        baseline_scores = self.baseline_scores
        score_payload: List[Dict[str, Any]] = []
        for score in scores:
            entry = {
                "supplier": score.supplier,
                "total_score": score.total_score,
                "metrics": score.metrics,
                "normalized_metrics": score.normalized_metrics,
                "rationale": score.rationale,
                "risk_flags": score.risk_flags,
            }
            base_value = baseline_scores.get(score.supplier)
            if base_value is not None:
                entry["baseline_total_score"] = base_value
                entry["score_delta"] = score.total_score - base_value
            score_payload.append(entry)
        payload = {
            "weights": normalized_weights,
            "scores": score_payload,
            "table": recalculated.reset_index(drop=True).to_dict(orient="records"),
        }
        self._what_if_history.append(payload)
        return payload

    @property
    def baseline_scores(self) -> Dict[str, float]:
        return {**self._latest_scores}

    def _load_file(self, file_obj: Any) -> pd.DataFrame:
        name = getattr(file_obj, "name", "")
        data = file_obj.read()
        file_obj.seek(0)
        stream = BytesIO(data)
        suffix = name.split(".")[-1].lower()
        if suffix in {"csv", "txt"}:
            df = pd.read_csv(stream)
        elif suffix in {"xls", "xlsx"}:
            df = pd.read_excel(stream)
        elif suffix in {"json"}:
            records = json.loads(data.decode("utf-8"))
            df = pd.DataFrame(records)
        elif suffix == "docx":
            df = _docx_to_dataframe(data)
        else:
            raise ValueError(f"Unsupported file format for {name}. Upload CSV, XLSX, JSON, or DOCX.")
        return df

    def _prepare_dataset(self, df: pd.DataFrame) -> pd.DataFrame:
        renamed = self._rename_columns(df)
        renamed = self._handle_missing(renamed)
        renamed = self._coerce_numeric(renamed)
        if not self.include_outliers:
            renamed = self._filter_outliers(renamed)
        return renamed

    def _rename_columns(self, df: pd.DataFrame) -> pd.DataFrame:
        lower_map = {col.lower(): col for col in df.columns}
        rename_rules = {}
        for canonical, aliases in COLUMN_ALIASES.items():
            for alias in aliases:
                if alias.lower() in lower_map:
                    rename_rules[lower_map[alias.lower()]] = canonical
                    break
        renamed = df.rename(columns=rename_rules)
        if "supplier" not in renamed.columns:
            raise ValueError("Supplier identifier column is required.")
        return renamed

    def _handle_missing(self, df: pd.DataFrame) -> pd.DataFrame:
        working = df.copy()
        numeric_cols = [col for col in DEFAULT_METRIC_WEIGHTS.keys() if col in working.columns]
        if self.missing_strategy == "drop":
            working = working.dropna(subset=numeric_cols)
            return working
        for col in numeric_cols:
            if working[col].isna().any():
                series = pd.to_numeric(working[col], errors="coerce")
                fill_value = median(series.dropna()) if not series.dropna().empty else 0
                working[col] = series.fillna(fill_value)
        working = working.fillna("")
        return working

    def _coerce_numeric(self, df: pd.DataFrame) -> pd.DataFrame:
        working = df.copy()
        for col in set(DEFAULT_METRIC_WEIGHTS).union({"risk", "sustainability", "hidden_costs", "lead_time", "price", "warranty", "quality", "compliance"}):
            if col in working.columns:
                working[col] = pd.to_numeric(working[col], errors="coerce")
        return working

    def _filter_outliers(self, df: pd.DataFrame) -> pd.DataFrame:
        working = df.copy()
        price_col = working.get("price")
        if price_col is None:
            return working
        series = pd.to_numeric(price_col, errors="coerce")
        z_scores = np.abs((series - series.mean()) / (series.std(ddof=0) or 1))
        filtered = working[z_scores < 3].copy()
        return filtered

    def _score_suppliers(
        self,
        df: pd.DataFrame,
        *,
        override_weights: Optional[Dict[str, float]] = None,
    ) -> tuple[pd.DataFrame, List[SupplierScore]]:
        working = df.copy()
        weights = override_weights or self.weight_config.normalized()
        metric_columns = [m for m in weights.keys() if m in working.columns]
        normalized_columns = {}
        for metric in metric_columns:
            normalized_columns[metric] = self._normalize_column(working[metric], higher_is_better=metric not in NEGATIVE_METRICS)
            working[f"normalized_{metric}"] = normalized_columns[metric]
        working["composite_score"] = sum(weights[m] * normalized_columns[m] for m in metric_columns)
        working = working.sort_values("composite_score", ascending=False).reset_index(drop=True)

        supplier_scores: List[SupplierScore] = []
        for _, row in working.iterrows():
            supplier = str(row.get("supplier", "Unknown Supplier"))
            metrics = {metric: float(row.get(metric, 0)) for metric in metric_columns}
            normalized = {metric: float(row.get(f"normalized_{metric}", 0)) for metric in metric_columns}
            rationale = self._compose_metric_rationale(metrics, normalized, weights)
            risk_flags = self._derive_risk_flags(row)
            supplier_scores.append(
                SupplierScore(
                    supplier=supplier,
                    total_score=float(row["composite_score"]),
                    metrics=metrics,
                    normalized_metrics=normalized,
                    rationale=rationale,
                    risk_flags=risk_flags,
                )
            )
        return working, supplier_scores

    @staticmethod
    def _normalize_column(series: pd.Series, *, higher_is_better: bool) -> pd.Series:
        numeric = pd.to_numeric(series, errors="coerce")
        if numeric.nunique(dropna=True) == 0:
            return pd.Series([0.5] * len(series))
        min_val = numeric.min()
        max_val = numeric.max()
        if pd.isna(min_val) or pd.isna(max_val) or max_val == min_val:
            normalized = pd.Series([0.5] * len(series))
        else:
            normalized = (numeric - min_val) / (max_val - min_val)
        normalized = normalized.clip(0, 1).fillna(0.5)
        if not higher_is_better:
            normalized = 1 - normalized
        return normalized

    def _compose_metric_rationale(
        self,
        metrics: Dict[str, float],
        normalized_metrics: Dict[str, float],
        weights: Dict[str, float],
    ) -> str:
        lines = []
        for metric, value in metrics.items():
            norm = normalized_metrics.get(metric, 0)
            weight = weights.get(metric, 0)
            descriptor = "strong" if norm >= 0.7 else "weak" if norm <= 0.3 else "moderate"
            lines.append(f"{metric.title()}: {value} (normalized {norm:.2f}, weight {weight:.2f}) — {descriptor}")
        return "\n".join(lines)

    def _derive_risk_flags(self, row: pd.Series) -> List[str]:
        flags = []
        if row.get("risk", 0) and row.get("risk", 0) > 7:
            flags.append("Elevated supplier risk score")
        if row.get("compliance", 0) and row.get("compliance", 0) < 5:
            flags.append("Compliance score below threshold")
        if row.get("hidden_costs") and row.get("hidden_costs") > 0:
            flags.append("Hidden costs identified")
        if self.rfp_anchor:
            target_price = float(self.rfp_anchor.get("target_price", 0) or 0)
            if target_price and row.get("price") and row.get("price") > target_price * 1.1:
                flags.append("Price exceeds RFP target by >10%")
        return flags

    def _derive_deltas(self, df: pd.DataFrame) -> List[str]:
        deltas = []
        if self.rfp_anchor:
            for key, target in self.rfp_anchor.items():
                if key in df.columns:
                    numeric_values = pd.to_numeric(df[key], errors="coerce")
                    median_value = numeric_values.median()
                    if isinstance(target, (int, float)):
                        delta = median_value - float(target)
                        if abs(delta) > 0.01:
                            deltas.append(f"Median {key} deviates from RFP anchor by {delta:.2f}.")
        top = df.head(1)
        if not top.empty:
            supplier = str(top.iloc[0].get("supplier", "Top Supplier"))
            score = float(top.iloc[0].get("composite_score", 0))
            deltas.append(f"{supplier} currently leads with composite score {score:.2f}.")
        if "hidden_costs" in df.columns and df["hidden_costs"].fillna(0).sum() > 0:
            deltas.append("Hidden costs detected across one or more bids.")
        return deltas

    def _llm_summary(
        self,
        df: pd.DataFrame,
        supplier_scores: List[SupplierScore],
        deltas: List[str],
    ) -> str:
        instructions = (
            "You are the Bid Comparison Strategy Agent. Produce an executive summary that:"
            "\n- Explains the ranking rationale"
            "\n- Highlights notable risks and compliance gaps"
            "\n- Suggests 1-2 negotiation levers based on the data"
            "\nRespond in Markdown with bullet points and a short concluding recommendation."
        )
        table_json = dataframe_to_pretty_json(df)
        score_snapshot = json.dumps([
            {
                "supplier": s.supplier,
                "total_score": round(s.total_score, 4),
                "top_flags": s.risk_flags[:3],
            }
            for s in supplier_scores
        ], indent=2)
        message = (
            f"Weighted metrics (normalized):\n```json\n{self.weight_config.normalized()}\n```\n"
            f"Supplier ranking details:\n```json\n{score_snapshot}\n```\n"
            f"Key deltas: {deltas}\n"
            f"Normalized table excerpt:\n```json\n{table_json}\n```"
        )
        return run_agent_sync("BidComparisonStrategist", instructions, message)


class MultiRFPComparisonEngine:
    def __init__(
        self,
        *,
        include_outliers: bool,
        missing_strategy: str,
        weight_config: WeightingConfig,
    ) -> None:
        self.include_outliers = include_outliers
        self.missing_strategy = missing_strategy
        self.weight_config = weight_config

    def analyse_vendor_across_rfps(
        self,
        vendor_file: Any,
        rfp_files: Sequence[Any],
    ) -> MultiRFPComparisonResult:
        from services.bid_analysis_service import BidComparisonEngine
        
        vendor_data = self._load_vendor_file(vendor_file)
        
        rfp_scores: List[RFPScenarioScore] = []
        scenario_rows: List[Dict[str, Any]] = []
        
        for rfp_file in rfp_files:
            rfp_name = getattr(rfp_file, "name", "RFP Scenario").rsplit(".", 1)[0]
            rfp_anchor = parse_rfp_spec(rfp_file)
            
            # Create a temporary engine with this RFP's anchor
            temp_engine = BidComparisonEngine(
                include_outliers=self.include_outliers,
                missing_strategy=self.missing_strategy,
                weight_config=self.weight_config,
                rfp_anchor=rfp_anchor,
            )
            
            # Score vendor against this RFP
            vendor_copy = vendor_data.copy()
            prepared = temp_engine._prepare_dataset(vendor_copy)
            scored, scores = temp_engine._score_suppliers(prepared)
            
            if scores:
                score_obj = scores[0]  # Single vendor
                fit_flags = self._derive_fit_flags(score_obj, rfp_anchor)
                rfp_score = RFPScenarioScore(
                    rfp_name=rfp_name,
                    total_score=score_obj.total_score,
                    metrics=score_obj.metrics,
                    normalized_metrics=score_obj.normalized_metrics,
                    rationale=score_obj.rationale,
                    fit_flags=fit_flags,
                )
                rfp_scores.append(rfp_score)
                
                row = {
                    "RFP Scenario": rfp_name,
                    "Composite Score": round(score_obj.total_score, 4),
                    **{f"{k}_normalized": round(v, 3) for k, v in score_obj.normalized_metrics.items()},
                }
                scenario_rows.append(row)
        
        scenario_table = pd.DataFrame(scenario_rows) if scenario_rows else pd.DataFrame()
        variation_notes = self._derive_variation_notes(rfp_scores)
        llm_summary = self._llm_summary(rfp_scores, variation_notes)
        
        return MultiRFPComparisonResult(
            scenario_table=scenario_table,
            weight_config=self.weight_config.normalized(),
            rfp_scores=rfp_scores,
            variation_notes=variation_notes,
            llm_summary=llm_summary,
        )

    def _load_vendor_file(self, file_obj: Any) -> pd.DataFrame:
        name = getattr(file_obj, "name", "")
        data = file_obj.read()
        file_obj.seek(0)
        stream = BytesIO(data)
        suffix = name.split(".")[-1].lower()
        if suffix in {"csv", "txt"}:
            df = pd.read_csv(stream)
        elif suffix in {"xls", "xlsx"}:
            df = pd.read_excel(stream)
        elif suffix in {"json"}:
            records = json.loads(data.decode("utf-8"))
            df = pd.DataFrame(records)
        elif suffix == "docx":
            df = _docx_to_dataframe(data)
        else:
            raise ValueError(f"Unsupported file format for {name}. Upload CSV, XLSX, JSON, or DOCX.")
        
        # Ensure supplier column exists
        if "supplier" not in df.columns:
            for col in SUPPLIER_ID_COLUMNS:
                if col in df.columns:
                    df = df.rename(columns={col: "supplier"})
                    break
            else:
                df["supplier"] = name.split(".")[0]
        
        return df

    def _derive_fit_flags(self, score: SupplierScore, rfp_anchor: Dict[str, Any]) -> List[str]:
        flags: List[str] = []
        target_price = rfp_anchor.get("target_price") or rfp_anchor.get("price")
        if target_price and isinstance(target_price, (int, float)):
            actual_price = score.metrics.get("price")
            if actual_price and actual_price > target_price * 1.1:
                flags.append(f"Price exceeds RFP target by {((actual_price / target_price - 1) * 100):.1f}%")
        
        if score.metrics.get("compliance", 0) < 7:
            flags.append("Compliance score below recommended threshold")
        
        if score.metrics.get("risk", 0) > 5:
            flags.append("Elevated risk profile for this RFP")
        
        return flags

    def _derive_variation_notes(self, rfp_scores: List[RFPScenarioScore]) -> List[str]:
        notes: List[str] = []
        if not rfp_scores:
            return notes
        
        scores = [s.total_score for s in rfp_scores]
        score_range = max(scores) - min(scores)
        if score_range > 0.15:
            notes.append(f"Significant score variation across RFPs (range: {score_range:.2f})")
        
        best = max(rfp_scores, key=lambda x: x.total_score)
        worst = min(rfp_scores, key=lambda x: x.total_score)
        notes.append(f"Best fit: {best.rfp_name} ({best.total_score:.2f})")
        notes.append(f"Weakest fit: {worst.rfp_name} ({worst.total_score:.2f})")
        
        all_flags = [flag for score in rfp_scores for flag in score.fit_flags]
        if all_flags:
            notes.append(f"Total fit flags raised: {len(all_flags)}")
        
        return notes

    def _llm_summary(self, rfp_scores: List[RFPScenarioScore], variation_notes: List[str]) -> str:
        instructions = (
            "You are the Multi-RFP Fit Analysis Agent. Produce an executive summary that:"
            "\n- Identifies which RFP scenarios the vendor is best suited for"
            "\n- Highlights gaps or weaknesses across different requirement contexts"
            "\n- Recommends which RFPs to prioritize or avoid"
            "\nRespond in Markdown with bullet points and a concluding recommendation."
        )
        
        score_snapshot = json.dumps([
            {
                "rfp": s.rfp_name,
                "total_score": round(s.total_score, 4),
                "fit_flags": s.fit_flags,
            }
            for s in rfp_scores
        ], indent=2)
        
        message = (
            f"Weighted metrics (normalized):\n```json\n{self.weight_config.normalized()}\n```\n"
            f"RFP scenario scores:\n```json\n{score_snapshot}\n```\n"
            f"Variation notes: {variation_notes}\n"
        )
        return run_agent_sync("MultiRFPFitAnalyst", instructions, message)


def parse_rfp_spec(uploaded_file: Any) -> dict:
    if not uploaded_file:
        return {}
    name = getattr(uploaded_file, "name", "")
    data = uploaded_file.read()
    uploaded_file.seek(0)
    suffix = name.split(".")[-1].lower()
    if suffix in {"json"}:
        raw = json.loads(data.decode("utf-8"))
        return {key: _coerce_anchor_value(value) for key, value in raw.items()}
    if suffix in {"csv"}:
        df = pd.read_csv(BytesIO(data))
        anchors: Dict[str, Any] = {}
        for row in df.values:
            if len(row) < 2:
                continue
            key = str(row[0]).strip()
            value = _coerce_anchor_value(row[1])
            anchors[key] = value
        return anchors
    if suffix in {"txt", "md"}:
        text = data.decode("utf-8")
        anchors = {}
        for line in text.splitlines():
            if ":" in line:
                key, value = line.split(":", 1)
                anchors[key.strip()] = _coerce_anchor_value(value)
        return anchors
    if suffix in {"docx"}:
        document = _load_docx_document(data)
        anchors_raw = _docx_to_key_values(document)
        if not anchors_raw:
            raise ValueError("DOCX spec must contain a table or key:value lines.")
        return {key.strip(): _coerce_anchor_value(val) for key, val in anchors_raw.items() if key}
    raise ValueError("Unsupported RFP spec format. Provide JSON, CSV, text, or DOCX.")

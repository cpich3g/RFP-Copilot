"""Streamlit page for negotiation strategy copilot."""
from __future__ import annotations

import json
import os
import re
import tempfile
import wave
from io import BytesIO
from typing import Any, Dict, List, Optional

import numpy as np
import streamlit as st
import azure.cognitiveservices.speech as speechsdk
from streamlit_webrtc import webrtc_streamer, WebRtcMode, AudioProcessorBase
import pandas as pd
from docx import Document

from components.settings_drawer import render_global_settings
from services.negotiation_service import (
    DEFAULT_TACTICS,
    NegotiationObjective,
    NegotiationStrategyEngine,
    PostNegotiationSummary,
    SupplierProfile,
)

st.set_page_config(page_title="Negotiation Strategy", layout="wide")


HISTORIC_SUPPLIERS: Dict[str, Dict[str, any]] = {
    "Northwind Traders": {
        "supplier_id": "Northwind Traders",
        "description": "Incumbent SaaS provider with strong support history.",
        "historic_notes": [
            "Delivered 98% on-time across the last 12 months",
            "Requested inflationary increase in the last renewal",
        ],
        "risk_level": "Medium",
        "last_negotiated": "2025-01-15",
        "category": "Cloud Services",
    },
    "Fourth Coffee": {
        "supplier_id": "Fourth Coffee",
        "description": "Regional logistics partner with responsive escalation team.",
        "historic_notes": [
            "Waived rush fees during Q2 peak",
            "Service credits issued for 2 delayed shipments",
        ],
        "risk_level": "Low",
        "last_negotiated": "2024-11-08",
        "category": "Logistics",
    },
    "Contoso Manufacturing": {
        "supplier_id": "Contoso Manufacturing",
        "description": "Strategic hardware supplier with mixed quality performance.",
        "historic_notes": [
            "Quality audit flagged soldering variance in 2023",
            "Offering bundled warranty program for new SKUs",
        ],
        "risk_level": "High",
        "last_negotiated": "2024-07-22",
        "category": "Hardware",
    },
}

PROFILE_FIELD_ALIASES: Dict[str, str] = {
    "supplier": "supplier_id",
    "supplier_id": "supplier_id",
    "supplier name": "supplier_id",
    "name": "supplier_id",
    "description": "description",
    "summary": "description",
    "overview": "description",
    "historic_notes": "historic_notes",
    "historic note": "historic_notes",
    "notes": "historic_notes",
    "risk": "risk_level",
    "risk level": "risk_level",
    "risk_rating": "risk_level",
    "risk rating": "risk_level",
    "last negotiated": "last_negotiated",
    "last_negotiated": "last_negotiated",
    "category": "category",
    "market": "category",
    "region": "region",
    "primary contact": "primary_contact",
    "contact": "primary_contact",
}


def _normalize_profile_key(raw_key: str) -> Optional[str]:
    key = raw_key.strip().lower().replace("-", " ")
    key = re.sub(r"\s+", " ", key)
    return PROFILE_FIELD_ALIASES.get(key, key.replace(" ", "_")) if key else None


def _parse_profile_value(field: str, value: str) -> Any:
    text = value.strip()
    if not text:
        return None
    if field == "historic_notes":
        parts = re.split(r"[;\n]", text)
        notes = []
        for part in parts:
            cleaned = part.strip().lstrip("•-·").strip()
            if cleaned:
                notes.append(cleaned)
        return notes
    return text


def _merge_profile_value(container: Dict[str, Any], field: str, value: Any) -> None:
    if value is None:
        return
    if field == "historic_notes":
        existing = container.get(field, [])
        if not isinstance(existing, list):
            existing = [existing]
        if isinstance(value, list):
            existing.extend(value)
        else:
            existing.append(value)
        # Deduplicate while preserving order
        seen = set()
        deduped = []
        for note in existing:
            if note not in seen:
                deduped.append(note)
                seen.add(note)
        container[field] = deduped
    else:
        container[field] = value


def _parse_supplier_profile_docx(file_bytes: bytes) -> Dict[str, Any]:
    document = Document(BytesIO(file_bytes))
    profile: Dict[str, Any] = {}

    def capture(raw_key: str, raw_value: str) -> None:
        normalized = _normalize_profile_key(raw_key)
        if not normalized:
            return
        parsed_value = _parse_profile_value(normalized, raw_value)
        _merge_profile_value(profile, normalized, parsed_value)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if len(cells) >= 2 and cells[0]:
                capture(cells[0], cells[1])

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if ":" in text:
            key, value = text.split(":", 1)
            capture(key, value)
        elif text.startswith(("•", "-")):
            note = text.lstrip("•- ").strip()
            if note:
                _merge_profile_value(profile, "historic_notes", note)

    return profile


def _init_state():
    st.session_state.setdefault("neg_engine", NegotiationStrategyEngine())
    st.session_state.setdefault("neg_profile", None)
    st.session_state.setdefault("neg_market", None)
    st.session_state.setdefault("neg_objectives", NegotiationObjective())
    st.session_state.setdefault("neg_transcript", [])
    st.session_state.setdefault("neg_voice_transcript", [])
    st.session_state.setdefault("neg_live_guidance", [])
    st.session_state.setdefault("neg_outcomes", [])
    st.session_state.setdefault("neg_voice_frames", [])
    st.session_state.setdefault("neg_voice_sample_rate", None)
    st.session_state.setdefault("neg_voice_channels", 1)
    st.session_state.setdefault("neg_uploaded_suppliers", {})
    st.session_state.setdefault("neg_selected_history", [])
    st.session_state.setdefault("neg_supplier_catalog", {})
    st.session_state.setdefault("neg_active_supplier_id", None)


def _settings_drawer():
    render_global_settings()
    st.sidebar.header("🎛️ Session Settings")
    st.sidebar.checkbox("Enable Voice Copilot", value=False, key="neg_voice_enabled")
    st.sidebar.checkbox("Auto-log transcript", value=True, key="neg_auto_log")
    st.sidebar.caption("Configure session defaults. Voice capture requires browser permissions.")


def _convert_frames_to_wav() -> Optional[BytesIO]:
    frames = st.session_state.get("neg_voice_frames", [])
    sample_rate = st.session_state.get("neg_voice_sample_rate") or 16000
    channels = st.session_state.get("neg_voice_channels", 1)
    if not frames:
        return None
    buffer = BytesIO()
    with wave.open(buffer, "wb") as wav_file:
        wav_file.setnchannels(channels)
        wav_file.setsampwidth(2)
        wav_file.setframerate(sample_rate)
        wav_file.writeframes(b"".join(frames))
    buffer.seek(0)
    st.session_state["neg_voice_frames"] = []
    return buffer


def _transcribe_voice_clip() -> Optional[str]:
    wav_buffer = _convert_frames_to_wav()
    if wav_buffer is None:
        st.warning("No voice audio captured yet.")
        return None
    speech_key = os.getenv("AZURE_SPEECH_KEY") or os.getenv("SPEECH_KEY")
    speech_region = os.getenv("AZURE_SPEECH_REGION") or os.getenv("SPEECH_REGION")
    if not speech_key or not speech_region:
        st.warning("Set AZURE_SPEECH_KEY and AZURE_SPEECH_REGION to enable transcription.")
        return None
    with tempfile.NamedTemporaryFile(delete=False, suffix=".wav") as tmp_file:
        tmp_file.write(wav_buffer.getvalue())
        tmp_path = tmp_file.name
    try:
        speech_config = speechsdk.SpeechConfig(subscription=speech_key, region=speech_region)
        audio_config = speechsdk.audio.AudioConfig(filename=tmp_path)
        recognizer = speechsdk.SpeechRecognizer(speech_config=speech_config, audio_config=audio_config)
        result = recognizer.recognize_once()
    finally:
        try:
            os.unlink(tmp_path)
        except OSError:
            pass
    if result.reason == speechsdk.ResultReason.RecognizedSpeech:
        return result.text
    if result.reason == speechsdk.ResultReason.NoMatch:
        st.warning("Speech service could not understand the audio clip.")
    elif result.reason == speechsdk.ResultReason.Canceled:
        details = result.cancellation_details if hasattr(result, "cancellation_details") else None
        message = details.error_details if details and details.error_details else "Unknown error"
        st.error(f"Speech transcription cancelled: {message}")
    return None


def _ingest_user_message(
    engine: NegotiationStrategyEngine,
    profile,
    objectives,
    market_context,
    message: str,
):
    if not message or not message.strip():
        return
    st.session_state["neg_transcript"].append({"role": "user", "content": message})
    response = engine.live_guidance(
        profile=profile,
        objectives=objectives,
        transcript=st.session_state["neg_transcript"],
        market_context=market_context,
    )
    st.session_state["neg_transcript"].append({"role": "assistant", "content": response})
    st.session_state["neg_live_guidance"].append(response)
    st.session_state.setdefault("neg_history", []).append({"prompt": message, "response": response})
    st.session_state["neg_voice_transcript"].append(message)


def _memory_to_tactics(memories: List[PostNegotiationSummary]) -> List[Dict[str, str]]:
    notes: List[Dict[str, str]] = []
    for summary in memories[-3:]:
        descriptor = summary.outcome or "Negotiation session"
        lessons = summary.lessons_learned or "Document post-session learnings to enrich this panel."
        notes.append(
            {
                "name": f"{descriptor[:40]}" + ("…" if len(descriptor) > 40 else ""),
                "description": lessons,
                "category": "Lessons",
            }
        )
    return notes


def _build_batna_notes(objectives: NegotiationObjective, market_context: Optional[Dict[str, Any]]) -> List[str]:
    notes: List[str] = []
    median_discount = (market_context or {}).get("median_discount")
    if isinstance(median_discount, (int, float)):
        notes.append(
            f"Anchor fallback offer at ~{median_discount:.0%} market discount to preserve leverage."
        )
    if objectives.target_price is not None:
        notes.append(
            f"Walk-away price set at {objectives.currency} {objectives.target_price:,.2f}."
        )
    if objectives.additional_constraints:
        notes.append("Non-negotiables: " + "; ".join(objectives.additional_constraints))
    if objectives.moq:
        notes.append(f"Minimum order quantity floor: {objectives.moq} units.")
    return notes


def _derive_risk_alerts(
    profile: Optional[SupplierProfile],
    market_context: Optional[Dict[str, Any]],
    objectives: NegotiationObjective,
) -> List[str]:
    alerts: List[str] = []
    if profile and profile.risk_level.lower() == "high":
        alerts.append("Supplier flagged high risk—prepare contingency clauses and tighter SLAs.")
    if profile and profile.historic_notes:
        alerts.append(f"Historic note: {profile.historic_notes[-1]}")
    sla_norm = (market_context or {}).get("sla_norm")
    if objectives.sla and sla_norm:
        alerts.append(f"Target SLA {objectives.sla} vs market norm {sla_norm}; justify any uplift with penalties.")
    if objectives.penalties:
        alerts.append("Penalties defined—ensure enforcement triggers are explicit in final draft.")
    context_currency = (market_context or {}).get("currency")
    if context_currency and objectives.currency != context_currency:
        alerts.append(
            f"Currency mismatch (target {objectives.currency} vs market {context_currency}); address FX exposure."
        )
    return alerts


def _build_counter_templates(objectives: NegotiationObjective) -> List[str]:
    templates: List[str] = []
    target_price = objectives.target_price
    templates.append(
        """
Counter Offer Package
- Price: {currency} {price}
- SLA: {sla}
- Warranty: {warranty}
- Penalties: {penalties}
- Extras: Service credits for misses beyond SLA window.
""".format(
            currency=objectives.currency,
            price=f"{target_price:,.2f}" if target_price is not None else "TBD",
            sla=objectives.sla or "Match incumbent",
            warranty=objectives.warranty or "Standard",
            penalties=objectives.penalties or "Escalating credits",
        )
    )
    templates.append(
        """
Escalation Counter Template
- Concede: Extend term to 24 months for price protection.
- Ask: Add performance bond and quarterly executive QBR.
- Risk Guard: Include claw-back on recurring revenue shortfall.
"""
    )
    return templates


def _ensure_system_message(
    profile: Optional[SupplierProfile],
    objectives: NegotiationObjective,
    market_context: Optional[Dict[str, Any]],
):
    if st.session_state.get("neg_transcript"):
        return
    supplier = profile.supplier_id if profile else "Supplier"
    summary_lines = [
        f"Negotiation with **{supplier}**.",
        f"Target price: {objectives.currency} {objectives.target_price:,.2f}" if objectives.target_price is not None else "Target price pending.",
        f"SLA goal: {objectives.sla or 'Follow incumbent SLA'}.",
    ]
    if market_context and market_context.get("recency"):
        summary_lines.append(f"Market benchmark recency: {market_context['recency']} ({market_context.get('notes', 'stub data')}).")
    st.session_state["neg_transcript"].append({
        "role": "system",
        "content": "  \n".join(summary_lines),
    })


def _render_tactics_panel(
    engine: NegotiationStrategyEngine,
    profile: Optional[SupplierProfile],
    objectives: NegotiationObjective,
    market_context: Optional[Dict[str, Any]],
):
    st.markdown("### Tactics & Alerts")
    memory_notes = _memory_to_tactics(engine.get_strategy_memory(profile.supplier_id)) if profile else []
    palette = engine.render_tactics_palette(additional_notes=memory_notes)
    st.markdown("#### Levers")
    for tactic in palette:
        st.markdown(f"- **{tactic['name']}** ({tactic['category']}): {tactic['description']}")

    batna_notes = _build_batna_notes(objectives, market_context)
    st.markdown("#### BATNA Notes")
    if batna_notes:
        for note in batna_notes:
            st.markdown(f"- {note}")
    else:
        st.caption("Set objectives to populate BATNA guidance.")

    risk_alerts = _derive_risk_alerts(profile, market_context, objectives)
    st.markdown("#### Risk Alerts")
    if risk_alerts:
        for alert in risk_alerts:
            st.markdown(f"- {alert}")
    else:
        st.caption("Risk alerts will surface once supplier context is loaded.")

    st.markdown("#### Counter-Proposal Templates")
    for template in _build_counter_templates(objectives):
        st.code(template.strip(), language="markdown")

class TranscriptAudioProcessor(AudioProcessorBase):
    def recv_audio(self, frame):
        audio = frame.to_ndarray()
        channels = 1
        if audio.ndim == 2:
            channels = audio.shape[0]
            audio = audio.T
        else:
            audio = audio.reshape(-1, 1)
        samples = audio.astype(np.float32)
        samples = np.clip(samples, -1.0, 1.0)
        pcm = (samples * 32767).astype(np.int16)
        interleaved = pcm.flatten().tobytes()
        st.session_state.setdefault("neg_voice_frames", []).append(interleaved)
        st.session_state["neg_voice_sample_rate"] = frame.sample_rate
        st.session_state["neg_voice_channels"] = channels
        return frame


def _render_voice_capture(
    engine: NegotiationStrategyEngine,
    profile,
    objectives,
    market_context,
):
    if not st.session_state.get("neg_voice_enabled"):
        return
    st.subheader("Voice Capture (Push-to-talk)")
    webrtc_ctx = webrtc_streamer(
        key="negotiation-voice",
        mode=WebRtcMode.SENDONLY,
        audio_processor_factory=TranscriptAudioProcessor,
        async_processing=True,
    )
    if webrtc_ctx and webrtc_ctx.state.playing:
        st.info("Recording… stop to process audio.")
    else:
        st.caption("Voice capture uses WebRTC with optional Azure Speech transcription.")
    if st.button("Transcribe Voice Input", key="transcribe_voice"):
        transcription = _transcribe_voice_clip()
        if transcription:
            st.success("Voice input transcribed.")
            _ingest_user_message(engine, profile, objectives, market_context, transcription)


def _supplier_selector() -> tuple[Optional[str], Dict[str, Dict[str, Any]]]:
    st.markdown("### Supplier Selection")
    history_options = list(HISTORIC_SUPPLIERS.keys())
    selected_history = st.multiselect(
        "Select supplier(s) from history",
        history_options,
        key="neg_history_selection",
    )
    st.session_state["neg_selected_history"] = selected_history

    catalog: Dict[str, Dict[str, Any]] = {supplier: HISTORIC_SUPPLIERS[supplier] for supplier in selected_history}

    uploaded_profiles = st.file_uploader(
        "Upload Supplier Profile(s) (JSON or DOCX)",
        type=["json", "docx"],
        accept_multiple_files=True,
    )
    if uploaded_profiles:
        for uploaded in uploaded_profiles:
            payload_bytes = uploaded.getvalue()
            suffix = uploaded.name.split(".")[-1].lower()
            data: Dict[str, Any]
            if suffix == "json":
                try:
                    data = json.loads(payload_bytes.decode("utf-8"))
                except json.JSONDecodeError as exc:
                    st.warning(f"Failed to parse {uploaded.name}: {exc}")
                    continue
            elif suffix == "docx":
                try:
                    data = _parse_supplier_profile_docx(payload_bytes)
                except Exception as exc:  # noqa: BLE001
                    st.warning(f"Failed to parse {uploaded.name}: {exc}")
                    continue
            else:
                st.warning(f"Unsupported supplier profile format: {uploaded.name}")
                continue

            supplier_id = (
                data.get("supplier_id")
                or data.get("supplier")
                or data.get("name")
                or uploaded.name.split(".")[0]
            )
            data.setdefault("supplier_id", supplier_id)
            if "historic_notes" in data and not isinstance(data["historic_notes"], list):
                value = str(data["historic_notes"])
                data["historic_notes"] = [note.strip() for note in value.split(";") if note.strip()]
            st.session_state.setdefault("neg_uploaded_suppliers", {})[supplier_id] = data
            catalog[supplier_id] = data
            uploaded.seek(0)

    catalog.update(st.session_state.get("neg_uploaded_suppliers", {}))
    st.session_state["neg_supplier_catalog"] = catalog

    if catalog:
        with st.expander("Supplier Cohort Overview", expanded=False):
            for supplier, payload in catalog.items():
                description = payload.get("description") or "--"
                last_negotiated = payload.get("last_negotiated") or "Unknown"
                st.markdown(f"- **{supplier}** — {description} _(Last negotiated {last_negotiated})_")

    active_choices = ["--"] + list(catalog.keys())
    active_supplier = st.selectbox("Active negotiation focus", active_choices, key="neg_active_supplier")
    if active_supplier == "--":
        active_supplier = None
    return active_supplier, catalog


def _render_prebrief(
    engine: NegotiationStrategyEngine,
    supplier_id: str,
    supplier_catalog: Dict[str, Dict[str, Any]],
):
    st.subheader("Pre-brief Builder")
    fallback_profile = {
        "supplier_id": supplier_id,
        "description": "Trusted supplier with solid regional presence.",
        "historic_notes": [
            "Delivered 98% on-time",
            "Upsell attempts in previous renewal",
        ],
        "risk_level": "Medium",
        "category": "Cloud Services",
    }
    supplier_profile_raw = supplier_catalog.get(supplier_id, fallback_profile)
    st.session_state["neg_profile_raw"] = supplier_profile_raw
    profile = engine.preload_supplier_profile(supplier_id, supplier_profile_raw)

    st.markdown("#### Supplier Snapshot")
    snap_cols = st.columns(3)
    snap_cols[0].metric("Risk Level", profile.risk_level)
    snap_cols[1].metric("Last Negotiated", profile.last_negotiated or "Unknown")
    snap_cols[2].metric("Historic Notes", len(profile.historic_notes))
    if profile.historic_notes:
        with st.expander("Historic Notes", expanded=False):
            for note in profile.historic_notes:
                st.markdown(f"- {note}")

    market_category = st.text_input(
        "Market Category",
        value=supplier_profile_raw.get("category", "Cloud Services"),
        key="neg_market_category",
    )
    if st.button("Fetch Market Benchmark", key="market_fetch"):
        st.session_state["neg_market"] = engine.fetch_market_benchmarks(market_category)
        st.success("Benchmark loaded.")
    market_context = st.session_state.get("neg_market") or engine.fetch_market_benchmarks(market_category)

    bench_cols = st.columns(4)
    bench_cols[0].metric("Median Discount", f"{market_context.get('median_discount', 0):.0%}")
    bench_cols[1].metric("SLA Norm", market_context.get("sla_norm", "--"))
    bench_cols[2].metric("Warranty Norm", market_context.get("warranty_norm", "--"))
    bench_cols[3].metric("Currency", market_context.get("currency", "--"))
    st.caption(f"Benchmark recency: {market_context.get('recency', 'N/A')} — {market_context.get('notes', 'Stub data')}.")

    existing_objectives = st.session_state.get("neg_objectives", NegotiationObjective())
    with st.form("objectives_form"):
        st.markdown("### Objectives & Constraints")
        obj_cols_left, obj_cols_right = st.columns(2)
        with obj_cols_left:
            target_price = st.number_input(
                "Target Price",
                min_value=0.0,
                value=float(existing_objectives.target_price or 0.0),
            )
            moq = st.number_input(
                "Minimum Order Quantity",
                min_value=0,
                value=int(existing_objectives.moq or 0),
            )
            currency = st.selectbox(
                "Currency",
                ["USD", "EUR", "GBP", "INR"],
                index=["USD", "EUR", "GBP", "INR"].index(existing_objectives.currency)
                if existing_objectives.currency in ["USD", "EUR", "GBP", "INR"]
                else 0,
            )
            incoterms = st.text_input("Incoterms", value=existing_objectives.incoterms or "FOB")
        with obj_cols_right:
            sla = st.text_input("SLA Target", value=existing_objectives.sla or market_context.get("sla_norm", "99.9% uptime"))
            warranty = st.text_input("Warranty Expectation", value=existing_objectives.warranty or market_context.get("warranty_norm", "24 months"))
            penalties = st.text_area(
                "Penalties/Clauses",
                value=existing_objectives.penalties or "Include SLA penalties for missed uptime.",
            )
            additional = st.text_area(
                "Additional Constraints",
                value="\n".join(existing_objectives.additional_constraints) or "Maintain annual price increase cap <3%.",
            )
        submitted = st.form_submit_button("Generate Pre-brief")

    objectives = NegotiationObjective(
        target_price=target_price,
        sla=sla,
        warranty=warranty,
        penalties=penalties,
        moq=int(moq) if moq else None,
        currency=currency,
        incoterms=incoterms,
        additional_constraints=[line.strip() for line in additional.splitlines() if line.strip()],
    )
    st.session_state["neg_objectives"] = objectives
    if submitted:
        prebrief = engine.build_prebrief(profile, market_context, objectives)
        st.session_state["neg_prebrief"] = prebrief
        st.session_state["neg_profile"] = profile
        st.session_state["neg_market"] = market_context
        st.session_state["neg_transcript"] = []
        _ensure_system_message(profile, objectives, market_context)
        st.success("Pre-brief generated.")

    if st.session_state.get("neg_prebrief"):
        st.markdown("### Pre-brief Summary")
        st.markdown(st.session_state["neg_prebrief"])


def _render_live_copilot(engine: NegotiationStrategyEngine):
    st.markdown("---")
    st.subheader("Live Copilot")
    profile = st.session_state.get("neg_profile")
    market_context = st.session_state.get("neg_market")
    objectives = st.session_state.get("neg_objectives")
    if not all([profile, market_context, objectives]):
        st.info("Generate a pre-brief first to unlock live copilot.")
        return
    _ensure_system_message(profile, objectives, market_context)

    col_main, col_side = st.columns([2.5, 1])
    with col_main:
        _render_voice_capture(engine, profile, objectives, market_context)
        transcript_container = st.container()
        with transcript_container:
            for message in st.session_state.get("neg_transcript", [])[-30:]:
                role = message.get("role", "assistant")
                role_key = role if role in {"user", "assistant", "system"} else "assistant"
                with st.chat_message(role_key):
                    st.markdown(message.get("content", ""))
        user_input = st.chat_input("Enter a note or quote from the supplier…")
        if user_input:
            _ingest_user_message(engine, profile, objectives, market_context, user_input)

    with col_side:
        _render_tactics_panel(engine, profile, objectives, market_context)


def _render_post_session(engine: NegotiationStrategyEngine, supplier_id: str):
    st.markdown("---")
    st.subheader("Post-Negotiation Logging")
    with st.form("post_session_form"):
        outcome = st.text_area("Outcome Summary", value="Agreement in principle with revised price.")
        concessions = st.text_area(
            "Concessions",
            value="Supplier agreed to 5% price reduction; client accepted 18-month term.",
        )
        actions = st.text_area("Next Steps", value="Draft contract amendment; schedule compliance review.")
        lessons = st.text_area("Lessons Learned", value="Highlight risk-sharing clauses earlier in discussion.")
        submitted = st.form_submit_button("Log Session")
    if submitted:
        summary = PostNegotiationSummary(
            outcome=outcome.strip(),
            concessions=[line.strip() for line in concessions.split(";") if line.strip()],
            action_items=[line.strip() for line in actions.split(";") if line.strip()],
            lessons_learned=lessons.strip(),
        )
        engine.log_outcome(supplier_id, summary)
        st.session_state.setdefault("neg_outcomes", []).append(summary)
        st.success("Session logged.")

    memory = engine.get_strategy_memory(supplier_id)
    if memory:
        st.markdown("### Strategy Memory")
        rows = [
            {
                "Logged": entry.timestamp.strftime("%Y-%m-%d %H:%M UTC"),
                "Outcome": entry.outcome,
                "Concessions": "; ".join(entry.concessions) or "--",
                "Next Steps": "; ".join(entry.action_items) or "--",
                "Lessons": entry.lessons_learned or "--",
            }
            for entry in memory[-10:]
        ]
        st.dataframe(pd.DataFrame(rows))
        st.markdown("#### Lessons Learned Digest")
        digest = "\n".join(f"- {entry.lessons_learned}" for entry in memory[-5:] if entry.lessons_learned)
        st.markdown(digest or "No lessons captured yet — log a session to build institutional knowledge.")


def main():
    _init_state()
    _settings_drawer()
    st.title("Negotiation Strategy Copilot")
    st.caption("Prepare, execute, and learn from supplier negotiations with AI-assisted guidance.")

    supplier_id, supplier_catalog = _supplier_selector()
    if supplier_id != st.session_state.get("neg_active_supplier_id"):
        st.session_state["neg_active_supplier_id"] = supplier_id
        st.session_state["neg_transcript"] = []
        st.session_state["neg_live_guidance"] = []
        st.session_state["neg_outcomes"] = []
        st.session_state["neg_prebrief"] = None
        st.session_state["neg_market"] = None
    if not supplier_id:
        st.info("Select at least one supplier and set an active focus to begin.")
        return

    engine: NegotiationStrategyEngine = st.session_state["neg_engine"]
    _render_prebrief(engine, supplier_id, supplier_catalog)
    _render_live_copilot(engine)
    _render_post_session(engine, supplier_id)


if __name__ == "__main__":
    main()
